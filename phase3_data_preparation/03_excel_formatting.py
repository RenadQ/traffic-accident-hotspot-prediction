import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Load your DataFrame
df = pd.read_excel("transformed_data_train_V88.xlsx")

# New output file name
output_path = "transformed_data_train_formatted.xlsx"
sheet_name = "Sheet1"

# Save initial Excel file
df.to_excel(output_path, sheet_name=sheet_name, index=False)

# Load workbook for formatting
wb = load_workbook(output_path)
ws = wb[sheet_name]

# Format header
header_fill = PatternFill(start_color="FFD966", end_color="FFD966", fill_type="solid")
header_font = Font(bold=True)
for cell in ws[1]:
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center")

# Apply zebra striping
zebra_fill = PatternFill(start_color="FFF8CC", end_color="FFF8CC", fill_type="solid")
for r in range(2, ws.max_row + 1):
    if r % 2 == 0:
        for c in range(1, ws.max_column + 1):
            ws.cell(row=r, column=c).fill = zebra_fill

# Apply thin border to all cells
thin = Side(style="thin", color="DDDDDD")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
for r in range(1, ws.max_row + 1):
    for c in range(1, ws.max_column + 1):
        ws.cell(row=r, column=c).border = border

# Auto-adjust column widths
for col in ws.columns:
    max_length = 0
    col_letter = get_column_letter(col[0].column)
    for cell in col:
        try:
            max_length = max(max_length, len(str(cell.value)))
        except:
            pass
    ws.column_dimensions[col_letter].width = min(max_length + 3, 40)

# Freeze header row
ws.freeze_panes = "A2"

# Save formatted workbook
wb.save(output_path)

print(f"✅ File '{output_path}' has been successfully saved and formatted.")
